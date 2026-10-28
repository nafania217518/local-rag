from pathlib import Path

import docx
import pymupdf

MIN_TEXT_CHARS = 20
RENDER_DPI = 150


def _nonempty(pages):
    return [(n, t) for n, t in pages if t.strip()]


def parse_pdf(path, cfg=None, parse_page=None):
    pages = []
    with pymupdf.open(path) as doc:
        for n, page in enumerate(doc, 1):
            text = page.get_text()
            if len(text.strip()) < MIN_TEXT_CHARS and parse_page:
                text = parse_page(cfg, page.get_pixmap(dpi=RENDER_DPI).tobytes("png"))
            pages.append((n, text))
    return _nonempty(pages)


def parse_docx(path, cfg=None, parse_page=None):
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    parts += [" | ".join(c.text for c in row.cells) for t in d.tables for row in t.rows]
    return _nonempty([(1, "\n\n".join(parts))])


def parse_txt(path, cfg=None, parse_page=None):
    return _nonempty([(1, Path(path).read_text(encoding="utf-8", errors="replace"))])


PARSERS = {".pdf": parse_pdf, ".docx": parse_docx, ".txt": parse_txt, ".md": parse_txt}


def parse_file(path, cfg=None, parse_page=None):
    suffix = Path(path).suffix.lower()
    if suffix not in PARSERS:
        raise ValueError(f"unsupported file type: {suffix}")
    return PARSERS[suffix](path, cfg, parse_page)
