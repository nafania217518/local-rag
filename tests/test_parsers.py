import docx
import pymupdf
import pytest

from localrag import parsers


def _pdf(path, pages):
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        if text:
            page.insert_text((72, 72), text)
        else:
            page.draw_rect(pymupdf.Rect(50, 50, 200, 200), color=(0, 0, 0), fill=(0, 0, 0))
    doc.save(path)
    return path


class FakeParse:
    def __init__(self):
        self.calls = []

    def __call__(self, cfg, png):
        self.calls.append(png)
        return "OCR TEXT"


def test_parse_pdf_extracts_text_per_page(tmp_path):
    p = _pdf(tmp_path / "a.pdf", ["Hello text page", "Second page sentence here"])
    assert [(n, t.strip()) for n, t in parsers.parse_pdf(p)] == [
        (1, "Hello text page"), (2, "Second page sentence here")]


def test_parse_pdf_routes_pages_without_text_to_parse_page(tmp_path):
    fake = FakeParse()
    p = _pdf(tmp_path / "s.pdf", [None])
    assert parsers.parse_pdf(p, {}, fake) == [(1, "OCR TEXT")]
    assert fake.calls[0].startswith(b"\x89PNG")


def test_parse_pdf_mixed_only_parses_image_pages(tmp_path):
    fake = FakeParse()
    p = _pdf(tmp_path / "m.pdf", ["Hello text page with enough chars", None])
    result = parsers.parse_pdf(p, {}, fake)
    assert len(fake.calls) == 1
    assert result[1] == (2, "OCR TEXT")
    assert "Hello text page" in result[0][1]


def test_parse_pdf_without_parse_page_skips_empty_pages(tmp_path):
    p = _pdf(tmp_path / "m.pdf", ["Hello text page", None])
    assert [n for n, _ in parsers.parse_pdf(p)] == [1]


def test_parse_docx_includes_paragraphs_and_table_cells(tmp_path):
    d = docx.Document()
    d.add_paragraph("First para")
    d.add_paragraph("Second para")
    t = d.add_table(rows=2, cols=2)
    for r in range(2):
        for c in range(2):
            t.cell(r, c).text = f"c{r}{c}"
    d.save(tmp_path / "a.docx")
    [(page, text)] = parsers.parse_docx(tmp_path / "a.docx")
    assert page == 1
    assert text == "First para\n\nSecond para\n\nc00 | c01\n\nc10 | c11"


def test_parse_txt_and_md(tmp_path):
    (tmp_path / "a.txt").write_text("plain", encoding="utf-8")
    (tmp_path / "b.md").write_text("# md", encoding="utf-8")
    assert parsers.parse_txt(tmp_path / "a.txt") == [(1, "plain")]
    assert parsers.parse_file(tmp_path / "b.md") == [(1, "# md")]


def test_parse_file_dispatches_by_suffix_case_insensitive(tmp_path):
    p = _pdf(tmp_path / "A.PDF", ["Hello text page"])
    assert parsers.parse_file(p)[0][1].strip() == "Hello text page"


def test_parse_file_rejects_unknown_suffix(tmp_path):
    with pytest.raises(ValueError, match="unsupported file type: .xyz"):
        parsers.parse_file(tmp_path / "a.xyz")


CLEAN_UK = (
    "Атомна енергетика України забезпечує понад половину виробництва електроенергії в країні. "
    "Підприємство експлуатує чотири атомні електростанції та постійно підвищує рівень безпеки. "
    "Персонал проходить регулярне навчання, а обладнання перевіряють незалежні експерти. "
    "Компанія також розвиває відновлювані джерела енергії, модернізує мережі та співпрацює "
    "з міжнародними партнерами для зміцнення енергетичної незалежності держави."
)
SOUP = " ".join(["ейИXX<ичок", "*Й)", "X<и*чочоX", "оHsX"] * 10 + ["атомна", "станція", "енергія"])


def test_text_quality_clean_page_scores_high():
    assert len(CLEAN_UK.split()) >= 40
    assert parsers.text_quality(CLEAN_UK) >= 0.95
    toc = "Терміни та визначення понять........................5\n" + CLEAN_UK
    assert parsers.text_quality(toc) >= 0.95


def test_text_quality_mixed_script_and_symbol_soup_scores_low():
    assert parsers.text_quality(SOUP) < 0.8


def test_text_quality_char_separated_text_capped():
    text = " ".join(["Е Н Е Р Г О А Т О М"] * 5 + ["атомна", "станція", "енергія", "безпека", "персонал"])
    assert parsers.text_quality(text) <= 0.5


def test_text_quality_short_text_is_neutral():
    assert parsers.text_quality("*Й) оHsX") == 1.0


def _pixmap():
    pm = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 50, 50), 0)
    pm.set_rect(pm.irect, (200, 50, 50))
    return pm


ASCII_SOUP = " ".join(["*X)<", "#$%a", "<<*>", "@!b?"] * 10)


def _scan_pdf(path, text, image=True):
    doc = pymupdf.open()
    page = doc.new_page()
    if image:
        page.insert_image(page.rect, pixmap=_pixmap())
    page.insert_text((72, 72), text)
    doc.save(path)
    return path


def test_is_scanned_requires_image_over_half_page(tmp_path):
    doc = pymupdf.open()
    doc.new_page().insert_image(doc[0].rect, pixmap=_pixmap())
    doc.new_page().insert_image(pymupdf.Rect(0, 0, 50, 50), pixmap=_pixmap())
    doc.new_page().insert_text((72, 72), "Text only page")
    doc.save(tmp_path / "i.pdf")
    with pymupdf.open(tmp_path / "i.pdf") as d:
        assert [parsers.is_scanned(p) for p in d] == [True, False, False]


def test_scanned_page_with_bad_text_layer_goes_to_ocr(tmp_path):
    p = _scan_pdf(tmp_path / "s.pdf", ASCII_SOUP)
    assert parsers.parse_pdf(p, {}, FakeParse()) == [(1, "OCR TEXT")]


def test_digital_page_with_bad_text_is_not_ocred(tmp_path):
    p = _scan_pdf(tmp_path / "d.pdf", ASCII_SOUP, image=False)
    fake = FakeParse()
    result = parsers.parse_pdf(p, {}, fake)
    assert fake.calls == [] and "*X)<" in result[0][1]


def test_force_ocr_reads_every_page(tmp_path):
    p = _pdf(tmp_path / "f.pdf", ["Hello text page one ok", "Hello text page two ok", "Hello text page three"])
    fake = FakeParse()
    result = parsers.parse_pdf(p, {}, fake, force_ocr=True)
    assert len(fake.calls) == 3 and [t for _, t in result] == ["OCR TEXT"] * 3


def test_on_page_reports_progress_and_ocr_flag(tmp_path):
    p = _pdf(tmp_path / "o.pdf", ["Hello text page with enough chars", None])
    calls = []
    parsers.parse_pdf(p, {}, FakeParse(), on_page=lambda *a: calls.append(a))
    assert calls == [(1, 2, False), (2, 2, True)]


def test_parse_file_forwards_keyword_arguments(tmp_path):
    p = _pdf(tmp_path / "k.pdf", ["Hello text page with enough chars"])
    fake = FakeParse()
    assert parsers.parse_file(p, {}, fake, force_ocr=True) == [(1, "OCR TEXT")]
    t = tmp_path / "k.txt"
    t.write_text("plain text")
    assert parsers.parse_file(t, {}, fake, force_ocr=True, on_page=print) == [(1, "plain text")]
